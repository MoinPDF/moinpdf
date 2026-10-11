"""Additional synthetic XLSX/PPTX fixtures. No user documents."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

out = Path.cwd() / 'fixtures'
out.mkdir(exist_ok=True)
geometry = out / 'geometry.png'
assert geometry.exists(), 'Run the project DOCX fixture generator first.'

workbook = Workbook()
sheet = workbook.active
sheet.title = 'Umsätze'
sheet.merge_cells('A1:D1')
sheet['A1'] = 'MoinPDF · XLSX Tabelle und Bild'
sheet['A1'].font = Font(name='Arial', size=18, bold=True, color='245D40')
sheet.row_dimensions[1].height = 32
sheet.append(['Artikel', 'Menge', 'Preis', 'Gesamt'])
for row in [['Äpfel', 3, 4.5, '=B3*C3'], ['Öl', 2, 19.9, '=B4*C4'], ['Straße', 1, 0, '=B5*C5']]:
    sheet.append(row)
sheet.append(['Summe', '', '', '=SUM(D3:D5)'])
for row in sheet.iter_rows(min_row=2, max_row=6, min_col=1, max_col=4):
    for cell in row:
        cell.border = Border(**{side: Side(style='thin', color='555555') for side in ('left', 'right', 'top', 'bottom')})
        cell.alignment = Alignment(vertical='center')
        cell.font = Font(name='Arial', size=12)
for cell in sheet[2]:
    cell.fill = PatternFill('solid', fgColor='245D40')
    cell.font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
for row in sheet.iter_rows(min_row=3, max_row=6, min_col=3, max_col=4):
    for cell in row:
        cell.number_format = '#,##0.00 "€"'
for column in ['A', 'B', 'C', 'D']:
    sheet.column_dimensions[column].width = 20
image = Image(str(geometry))
image.width, image.height = 360, 120
sheet.add_image(image, 'A9')
chart = BarChart()
chart.title = 'Mengenvergleich'
chart.add_data(Reference(sheet, min_col=2, min_row=2, max_row=5), titles_from_data=True)
chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=5))
chart.width, chart.height = 15, 7
sheet.add_chart(chart, 'A17')
sheet.print_area = 'A1:F32'
sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
sheet.page_setup.orientation = 'portrait'
sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
sheet.sheet_properties.pageSetUpPr.fitToPage = True
sheet.oddFooter.center.text = 'MoinPDF · Seite &P'
workbook.save(out / 'rich.xlsx')

presentation = Presentation()
presentation.slide_width, presentation.slide_height = Inches(10), Inches(5.625)
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
box = slide.shapes.add_textbox(Inches(.5), Inches(.35), Inches(9), Inches(.7))
paragraph = box.text_frame.paragraphs[0]
paragraph.text = 'MoinPDF · PPTX Tabelle und Bild'
paragraph.font.name = 'Arial'
paragraph.font.size, paragraph.font.bold = Pt(26), True
paragraph.font.color.rgb = RGBColor.from_string('245D40')
table = slide.shapes.add_table(4, 3, Inches(.6), Inches(1.35), Inches(8.8), Inches(1.65)).table
for i, row in enumerate([['Artikel', 'Menge', 'Preis'], ['Äpfel', '3', '4,50 €'], ['Öl', '2', '19,90 €'], ['Straße', '1', '0,00 €']]):
    for j, text in enumerate(row):
        cell = table.cell(i, j)
        cell.text = text
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.name, paragraph.font.size = 'Arial', Pt(17)
slide.shapes.add_picture(str(geometry), Inches(.7), Inches(3.35), width=Inches(4.5))
box = slide.shapes.add_textbox(Inches(5.3), Inches(3.6), Inches(4), Inches(1))
paragraph = box.text_frame.paragraphs[0]
paragraph.text = 'Grüße, ÄÖÜ äöü ß, Café, € § → α Ω.'
paragraph.font.name, paragraph.font.size = 'Arial', Pt(20)
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
box = slide.shapes.add_textbox(Inches(.6), Inches(.6), Inches(8.8), Inches(4))
frame = box.text_frame
frame.text = 'Seite 2 · ENDE DER PRÜFDATEI'
frame.paragraphs[0].font.size = Pt(30)
for text in ['Griechisch: Καλημέρα κόσμε.', 'Kyrillisch: Привет мир.']:
    paragraph = frame.add_paragraph()
    paragraph.text, paragraph.font.size = text, Pt(24)
presentation.save(out / 'rich.pptx')
print('created rich.xlsx rich.pptx')
