"""Generate a synthetic Office acceptance fixture. No personal documents."""
from pathlib import Path
import sys
from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw
ROOT=Path(sys.argv[1]).resolve()
ROOT.mkdir(parents=True,exist_ok=True)

image=Image.new('RGB',(480,160),'white');draw=ImageDraw.Draw(image)
draw.rectangle((15,15,150,145),fill='#2F7AB5');draw.ellipse((175,15,305,145),fill='#E69138');draw.polygon([(335,145),(400,15),(465,145)],fill='#38761D')
image.save(ROOT/'geometry.png')
d=Document();sec=d.sections[0];sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(0.75)
d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(11)
sec.header.paragraphs[0].text='MoinPDF · Prüfdokument'
sec.footer.paragraphs[0].text='Fußzeile – vertraulich'
d.add_heading('MoinPDF: DOCX → PDF',0)
p=d.add_paragraph();p.add_run('Fett: ').bold=True;p.add_run('Äpfel, Öl, Größe, Straße – Café; 19,90 €').bold=True
p=d.add_paragraph();p.add_run('Kursiv: résumé, naïve, déjà-vu. ').italic=True;p.add_run('Unterstrichen.').underline=True
d.add_paragraph('Ein Absatz mit vollständigen deutschen Sätzen. Der Zeilenumbruch muss lesbar bleiben, und Zahlen wie 123.456,78 sowie Sonderzeichen §, %, &, @ und © dürfen nicht verloren gehen. '*3)
d.add_heading('Tabelle und Bild',1)
table=d.add_table(rows=1,cols=3);table.style='Table Grid'
for c,t in zip(table.rows[0].cells,['Artikel','Menge','Preis']):c.text=t
for vals in [('Äpfel','3','4,50 €'),('Öl','2','19,90 €'),('Straße','1','0,00 €')]:
    for c,t in zip(table.add_row().cells,vals):c.text=t
d.add_picture(str(ROOT/'geometry.png'),width=Inches(3))
d.add_paragraph('Bildunterschrift: blaues Quadrat, orangefarbener Kreis, grünes Dreieck.')
d.add_page_break()
d.add_heading('Seite 2 – Schriften und Unicode',1)
for font in ['Arial','Calibri','Times New Roman','Courier New','Cambria']:
    p=d.add_paragraph();r=p.add_run(font+': Grüße, ÄÖÜ äöü ß, Café, € § → α Ω.');r.font.name=font;r.font.size=Pt(12)
for text in ['Griechisch: Καλημέρα κόσμε.','Kyrillisch: Привет мир.','Japanisch: 日本語の文書です。']:
    d.add_paragraph(text)
d.add_paragraph('ENDE DER PRÜFDATEI')
d.save(ROOT/'rich.docx')
print(ROOT/'rich.docx')
