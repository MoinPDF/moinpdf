import base64
import io
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from fastapi.testclient import TestClient
from server import app as service
from server.accessible_documents import create_document, validate_request
from server.test_service import pdf_bytes


def request():
    from PIL import Image
    buffer = io.BytesIO()
    Image.new('RGB', (80, 40), (10, 100, 160)).save(buffer, format='PNG')
    return {'title': 'MoinPDF Barrierefreiheit', 'language': 'de-DE', 'blocks': [
        {'kind': 'h1', 'text': 'Jahresbericht 2026'},
        {'kind': 'p', 'text': 'Erster Absatz mit Umlauten: ÄÖÜ. Preis: 123 Euro.'},
        {'kind': 'h2', 'text': 'Ergebnisse'},
        {'kind': 'figure', 'alt': 'Blaues Rechteck als Testabbildung.', 'data': base64.b64encode(buffer.getvalue()).decode()},
        {'kind': 'p', 'text': 'Letzter Absatz in der festgelegten Lesereihenfolge.'},
    ]}


class AccessibleTests(unittest.TestCase):
    def setUp(self):
        self.environment = patch.dict(os.environ, {'MOINPDF_API_KEY': 'synthetic-accessibility-' + 'x' * 32})
        self.environment.start()
        self.client = TestClient(service.app)
        self.headers = {'Authorization': 'Bearer ' + os.environ['MOINPDF_API_KEY']}

    def tearDown(self):
        self.client.close()
        self.environment.stop()

    def post(self, value):
        return self.client.post('/v1/pdfua/create', headers=self.headers, files={'request': ('struktur.json', json.dumps(value), 'application/json')})

    def test_reject_invalid_structure_and_unsafe_images(self):
        for change in [{'title': ''}, {'language': 'invalid'}, {'blocks': [{'kind': 'h3', 'text': 'Skipped'}]},
                       {'blocks': [{'kind': 'figure', 'alt': '', 'data': 'AAAA'}]},
                       {'blocks': [{'kind': 'figure', 'alt': 'Remote', 'data': 'https://example.com'}]},
                       {'blocks': [{'kind': 'p', 'text': 'x\x00x'}]},
                       {'blocks': [{'kind': 'p', 'text': 'x', 'unexpected': 1}]}]:
            value = {**request(), **change}
            with self.assertRaises(ValueError): validate_request(value)
            self.assertEqual(self.post(value).status_code, 422)

    def test_docx_contains_order_language_and_alt(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'out.docx'
            create_document(request(), path)
            document = Document(path)
            self.assertEqual(document.core_properties.title, request()['title'])
            self.assertEqual(document.core_properties.language, 'de-DE')
            self.assertEqual(document.paragraphs[0].style.name, 'Heading 1')
            self.assertIn('Letzter Absatz', document.paragraphs[-1].text)
            self.assertEqual(document.inline_shapes[0]._inline.docPr.get('descr'), 'Blaues Rechteck als Testabbildung.')

    def test_real_pdfua_export_and_independent_readback(self):
        import pymupdf
        response = self.post(request())
        self.assertEqual(response.status_code, 200, response.text[:500])
        with zipfile.ZipFile(io.BytesIO(response.content)) as bundle:
            report = json.loads(bundle.read('Ergebnis.json'))
            self.assertTrue(report['technicalPassed'], bundle.read('veraPDF-Bericht.xml').decode()[:12000])
            self.assertTrue(report['humanChecksRequired'])
            output = bundle.read('PDF-UA-1-technisch-geprueft.pdf')
        with pymupdf.open(stream=output, filetype='pdf') as pdf:
            content = ''.join(p.get_text() for p in pdf)
            self.assertIn('Jahresbericht 2026', content)
            self.assertIn('Letzter Absatz', content)
            self.assertEqual(pdf.metadata['title'], request()['title'])
            self.assertNotEqual(pdf.xref_get_key(pdf.pdf_catalog(), 'StructTreeRoot')[0], 'null')
            self.assertIn('pdfuaid:part', pdf.get_xml_metadata())
            objects = '\n'.join(pdf.xref_object(i) for i in range(1, pdf.xref_length()))
            self.assertIn('/H1', objects)
            self.assertIn('/Figure', objects)
            self.assertIn('/Alt', objects)
        checked = self.client.post('/v1/pdfua/validate', headers=self.headers, files={'file': ('out.pdf', output, 'application/pdf')}, data={'profile': 'ua1'})
        self.assertEqual(checked.status_code, 200, checked.text[:500])
        with zipfile.ZipFile(io.BytesIO(checked.content)) as bundle:
            self.assertTrue(json.loads(bundle.read('Ergebnis.json'))['technicalPassed'])

    def test_untagged_pdf_fails_ua1_and_ua2(self):
        for profile in ['ua1', 'ua2']:
            response = self.client.post('/v1/pdfua/validate', headers=self.headers, files={'file': ('raw.pdf', pdf_bytes(), 'application/pdf')}, data={'profile': profile})
            self.assertEqual(response.status_code, 200, response.text[:500])
            with zipfile.ZipFile(io.BytesIO(response.content)) as bundle:
                self.assertFalse(json.loads(bundle.read('Ergebnis.json'))['technicalPassed'])
                self.assertFalse(any(name.endswith('.pdf') for name in bundle.namelist()))

    def test_failed_validation_never_returns_pdf(self):
        with patch.object(service, 'pdfua_report', return_value=({'technicalPassed': False, 'note': 'Failed', 'humanChecksRequired': True}, b'<failed/>')):
            response = self.post(request())
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(response.content)) as bundle:
            self.assertFalse(any(name.endswith('.pdf') for name in bundle.namelist()))
        self.assertIn('fehlgeschlagen', response.headers['content-disposition'])

    def test_incomplete_or_wrong_validation_report_fails_closed(self):
        for raw in [b'<report/>', b'<report><validationReport profileName="PDF/A-2b" isCompliant="true"/><batchSummary totalJobs="1"/></report>',
                    b'<report><validationReport profileName="PDF/UA-1" isCompliant="true"/><batchSummary totalJobs="1" outOfMemory="1"/></report>']:
            with tempfile.TemporaryDirectory() as d, patch.object(service, 'command', return_value=(0, raw)):
                with self.assertRaises(service.HTTPException): service.pdfua_report(Path(d)/'x.pdf', 'ua1', Path(d))


if __name__ == '__main__': unittest.main()
