"""Bounded structured-document input for tagged, independently validated PDF/UA export."""
from __future__ import annotations
import base64
import io
import json
import re
import sys
from pathlib import Path

LANGUAGES = {'de-DE', 'en-US', 'fr-FR', 'es-ES', 'it-IT', 'nl-NL', 'pl-PL'}
MAX_REQUEST = 8 * 1024 * 1024


def text(value, maximum, label):
    if not isinstance(value, str) or not value.strip() or len(value) > maximum or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', value):
        raise ValueError(f'{label}: Bitte einen nicht leeren Text mit höchstens {maximum} Zeichen verwenden.')
    return value.strip()


def validate_request(value):
    if not isinstance(value, dict) or set(value) != {'title', 'language', 'blocks'}:
        raise ValueError('Ungültiger Strukturauftrag.')
    title = text(value['title'], 300, 'Dokumenttitel')
    if value['language'] not in LANGUAGES:
        raise ValueError('Bitte eine unterstützte Dokumentsprache wählen.')
    blocks = value['blocks']
    if not isinstance(blocks, list) or not 1 <= len(blocks) <= 200:
        raise ValueError('Bitte 1 bis 200 Inhaltsblöcke verwenden.')
    validated, previous_heading, characters, figures = [], 0, 0, 0
    for block in blocks:
        if not isinstance(block, dict):
            raise ValueError('Ungültiger Inhaltsblock.')
        kind = block.get('kind')
        if kind in {'h1', 'h2', 'h3', 'p'}:
            if set(block) != {'kind', 'text'}:
                raise ValueError('Ungültiger Textblock.')
            content = text(block['text'], 10000, 'Inhaltsblock')
            characters += len(content)
            if kind != 'p':
                level = int(kind[1])
                if level > previous_heading + 1:
                    raise ValueError('Überschriftenebenen nicht überspringen: mit Überschrift 1 beginnen.')
                previous_heading = level
            validated.append({'kind': kind, 'text': content})
        elif kind == 'figure':
            if set(block) != {'kind', 'alt', 'data'}:
                raise ValueError('Ungültiger Bildblock.')
            alt = text(block['alt'], 2000, 'Bildbeschreibung')
            characters += len(alt)
            figures += 1
            if figures > 10 or not isinstance(block['data'], str) or len(block['data']) > 4 * 1024 * 1024:
                raise ValueError('Höchstens zehn Bilder mit je 3 MB verwenden.')
            try:
                from PIL import Image
                data = base64.b64decode(block['data'], validate=True)
                with Image.open(io.BytesIO(data)) as image:
                    if image.format not in {'PNG', 'JPEG'} or not 1 <= image.width * image.height <= 4_000_000:
                        raise ValueError()
                    width, height = image.size
                    image.verify()
            except Exception:
                raise ValueError('Bilder müssen gültige PNG- oder JPEG-Dateien mit höchstens 4 Megapixeln sein.')
            validated.append({'kind': kind, 'alt': alt, 'bytes': data, 'width': width, 'height': height})
        else:
            raise ValueError('Erlaubt sind Überschriften, Absätze und Bilder.')
    if characters > 100000:
        raise ValueError('Höchstens 100.000 Zeichen pro Dokument verwenden.')
    return {'title': title, 'language': value['language'], 'blocks': validated}


def create_document(request, output: Path):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    settings = validate_request(request)
    document = Document()
    document.core_properties.title = settings['title']
    document.core_properties.language = settings['language']
    document.core_properties.author = 'MoinPDF'
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(2)
    for name, size in [('Normal', 12), ('Heading 1', 20), ('Heading 2', 16), ('Heading 3', 14)]:
        style = document.styles[name]
        style.font.name = 'Liberation Sans'
        style.font.size = Pt(size)
        lang = OxmlElement('w:lang')
        lang.set(qn('w:val'), settings['language'])
        style.element.get_or_add_rPr().append(lang)
        style.paragraph_format.space_after = Pt(8)
    for block in settings['blocks']:
        kind = block['kind']
        if kind == 'figure':
            # Fit in printable page, retaining aspect ratio. Images have a real
            # alternative description on the drawing object, not a visible caption.
            ratio = block['width'] / block['height']
            width_cm = min(17, 20 * ratio)
            shape = document.add_picture(io.BytesIO(block['bytes']), width=Cm(width_cm))
            shape._inline.docPr.set('descr', block['alt'])
            shape._inline.docPr.set('title', '')
        elif kind == 'p':
            document.add_paragraph(block['text'])
        else:
            document.add_heading(block['text'], int(kind[1]))
    document.save(output)


if __name__ == '__main__':
    directory = Path(sys.argv[1])
    try:
        raw = (directory / 'request.json').read_bytes()
        if len(raw) > MAX_REQUEST:
            raise ValueError('Der Strukturauftrag ist größer als 8 MB.')
        create_document(json.loads(raw), directory / 'structured.docx')
    except (ValueError, TypeError, KeyError):
        # No user-supplied document text or internal paths in error responses.
        (directory / 'structure-error.txt').write_text('Strukturauftrag ungültig. Titel, Sprache, Überschriftenfolge, Bildbeschreibungen und Größen prüfen.')
        sys.exit(2)
