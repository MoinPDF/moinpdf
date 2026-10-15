import io
import json
import os
import tempfile
import threading
import unittest
import zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from unittest.mock import patch

import cv2
import numpy as np
import pymupdf as fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from fastapi.testclient import TestClient
from server import app as service
from server.extended_documents import curve_image


def fixture():
    document=fitz.open();page=document.new_page()
    page.insert_text((40,60),'Dokumentenpruefung: Beispiel',fontsize=16)
    page.draw_rect((40,100,220,120));page.draw_rect((40,150,54,164));page.draw_line((40,210),(250,210))
    page.insert_text((40,300),'HiddenName',fontsize=14)
    hidden=page.search_for('HiddenName')[0];page.draw_rect(hidden,fill=(0,0,0),color=None)
    picture=np.full((40,60,3),(40,130,200),dtype=np.uint8)
    page.insert_image(fitz.Rect(40,400,240,530),stream=cv2.imencode('.png',picture)[1].tobytes())
    data=document.tobytes();document.close();return data


class ExtendedTests(unittest.TestCase):
    def setUp(self):
        self.environment=patch.dict(os.environ,{'MOINPDF_API_KEY':'synthetic-'+40*'x'});self.environment.start()
        self.client=TestClient(service.app);self.headers={'Authorization':'Bearer '+os.environ['MOINPDF_API_KEY']};self.pdf=fixture()

    def tearDown(self):
        self.client.close();self.environment.stop()

    def post(self,action,data=None,options=None):
        return self.client.post('/v1/documents/'+action,headers=self.headers,files={'file':('input.html' if action=='html' else 'input.pdf',data or self.pdf)},data={'options':json.dumps(options or {})})

    def test_powerpoint_and_html(self):
        response=self.post('powerpoint');self.assertEqual(response.status_code,200,response.text[:200])
        deck=Presentation(io.BytesIO(response.content));self.assertEqual(len(deck.slides),1);self.assertEqual(len(deck.slides[0].shapes),1)
        self.assertGreater(len(deck.slides[0].shapes[0].image.blob),100)
        response=self.post('html',b'<h1>Webseitenexport</h1><script>PRIVATE_SCRIPT</script><img src="file:///etc/passwd"><table><tr><td>Tabellenwert</td><td>12 Euro</td></tr></table>')
        self.assertEqual(response.status_code,200,response.text[:200])
        with fitz.open(stream=response.content,filetype='pdf') as pdf:text=''.join(p.get_text() for p in pdf)
        self.assertIn('Webseitenexport',text);self.assertIn('Tabellenwert',text);self.assertNotIn('PRIVATE_SCRIPT',text);self.assertNotIn('root:',text)

    def test_detected_form_is_editable_and_retains_page(self):
        response=self.post('form');self.assertEqual(response.status_code,200,response.text[:200])
        with fitz.open(stream=response.content,filetype='pdf') as pdf:
            page=pdf[0];widgets=list(page.widgets());self.assertGreaterEqual(len(widgets),3)
            self.assertIn(fitz.PDF_WIDGET_TYPE_CHECKBOX,[w.field_type for w in widgets])
            text=next(w for w in widgets if w.field_type==fitz.PDF_WIDGET_TYPE_TEXT);text.field_value='Testwert';text.update()
            self.assertIn('Dokumentenpruefung',pdf[0].get_text())
            saved=pdf.tobytes()
        with fitz.open(stream=saved,filetype='pdf') as pdf:self.assertIn('Testwert',[w.field_value for w in pdf[0].widgets()])

    def test_checks_find_real_problems(self):
        response=self.post('print',options={'dpi':150});self.assertEqual(response.status_code,200,response.text[:200]);data=response.json()
        self.assertTrue(any('Bildauflösung' in line for item in data['items'] for line in item.get('findings',[])))
        response=self.post('accessibility');self.assertEqual(response.status_code,200,response.text[:200])
        self.assertFalse(response.json()['items'][0]['structure_tree_present'])
        response=self.post('redaction',options={'terms':'HiddenName'});self.assertEqual(response.status_code,200,response.text[:200])
        self.assertIn('Text auslesbar',[item['status'] for item in response.json()['items']])
        self.assertIn('Verdächtige Überdeckung',[item['status'] for item in response.json()['items']])

    def test_book_curve_correction(self):
        pixels=np.full((180,220,3),255,dtype=np.uint8)
        for x in range(220):
            y=round(65+20*(1-x/219)**2);pixels[y-1:y+2,x]=0
        corrected=curve_image(pixels,20,'left');positions=[]
        for x in range(5,215):positions.append(np.where(corrected[:,x,0]<100)[0].mean())
        self.assertLess(np.std(positions),.65)
        response=self.post('dewarp',options={'curve':5,'side':'left'});self.assertEqual(response.status_code,200,response.text[:200])
        with fitz.open(stream=response.content,filetype='pdf') as pdf:self.assertEqual(len(pdf),1);self.assertEqual(len(pdf[0].get_images()),1)

    def test_orientation_real_tesseract(self):
        document=fitz.open();page=document.new_page()
        lines='\n'.join('This is an example sentence for automatic page orientation.' for _ in range(32))
        page.insert_textbox(fitz.Rect(35,35,560,805),lines,fontsize=14,fontname='tiro');page.set_rotation(180)
        response=self.post('orient',document.tobytes(),{'confidence':5});document.close()
        self.assertEqual(response.status_code,200,response.text[:200]);archive=zipfile.ZipFile(io.BytesIO(response.content))
        report=json.loads(archive.read('Pruefbericht.json'));self.assertEqual(report['items'][0]['rotation'],180,report)
        with fitz.open(stream=archive.read('Ausgerichtet.pdf'),filetype='pdf') as pdf:self.assertEqual(pdf[0].rotation,0);self.assertIn('example',pdf[0].get_text())

    def test_word_excel_merge_across_formatting_runs(self):
        doc=Document();paragraph=doc.add_paragraph();paragraph.add_run('Hallo {{Na').bold=True;paragraph.add_run('me}}, aus {{Ort}}!')
        doc.sections[0].header.paragraphs[0].text='Brief an {{Name}}'
        template=io.BytesIO();doc.save(template)
        workbook=Workbook();sheet=workbook.active;sheet.append(['Name','Ort']);sheet.append(['Anna','Köln']);sheet.append(['Jörg','Bremen']);data=io.BytesIO();workbook.save(data)
        response=self.client.post('/v1/word-merge',headers=self.headers,files={'file':('template.docx',template.getvalue()),'spreadsheet':('data.xlsx',data.getvalue())})
        self.assertEqual(response.status_code,200,response.text[:200]);archive=zipfile.ZipFile(io.BytesIO(response.content))
        for index,(name,city) in enumerate([('Anna','Köln'),('Jörg','Bremen')],1):
            with fitz.open(stream=archive.read(f'Serienbrief-{index:03}.pdf'),filetype='pdf') as pdf:text=''.join(p.get_text() for p in pdf)
            self.assertIn(name,text);self.assertIn(city,text);self.assertNotIn('{{',text)

    def test_translation_http_contract_and_side_by_side_output(self):
        calls=[]
        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                payload=json.loads(self.rfile.read(int(self.headers['Content-Length'])));calls.append(payload)
                body=json.dumps({'translatedText':'Uebersetzter Testtext'}).encode();self.send_response(200);self.send_header('Content-Type','application/json');self.end_headers();self.wfile.write(body)
            def log_message(self,*args):pass
        server=HTTPServer(('127.0.0.1',0),Handler);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        try:
            with patch.dict(os.environ,{'MOINPDF_TRANSLATE_URL':f'http://127.0.0.1:{server.server_port}'}):response=self.post('translate',options={'target':'de','language':'en'})
            self.assertEqual(response.status_code,200,response.text[:200]);self.assertEqual(calls[0]['target'],'de')
            with fitz.open(stream=response.content,filetype='pdf') as pdf:text=''.join(p.get_text() for p in pdf)
            self.assertIn('Uebersetzter Testtext',text);self.assertIn('Dokumentenpruefung',text)
            self.assertIn('Original und Übersetzung',text)
            with patch.dict(os.environ,{'MOINPDF_TRANSLATE_URL':f'http://127.0.0.1:{server.server_port}'}):english=self.post('translate',options={'target':'de','language':'en','uiLanguage':'en'})
            self.assertEqual(english.status_code,200,english.text[:200])
            with fitz.open(stream=english.content,filetype='pdf') as pdf:english_text=''.join(p.get_text() for p in pdf)
            self.assertIn('Original and translation',english_text);self.assertIn('Page 1',english_text)
            self.assertIn('Uebersetzter Testtext',english_text);self.assertIn('Dokumentenpruefung',english_text)
        finally:server.shutdown();server.server_close();thread.join()

    @unittest.skipUnless(service.MUSTANG.is_file(),'Mustang nicht installiert')
    def test_zugferd_creation_and_invalid_input_reporting(self):
        xml=Path('tests/fixtures/invoice-cii.xml').read_bytes()
        response=self.client.post('/v1/zugferd/create',headers=self.headers,files={'file':('invoice.pdf',self.pdf),'invoice':('invoice.xml',xml)},data={'profile':'X'})
        self.assertEqual(response.status_code,200,response.text[:200]);archive=zipfile.ZipFile(io.BytesIO(response.content))
        summary=json.loads(archive.read('Ergebnis.json'));self.assertTrue(summary['compliant'],summary)
        with fitz.open(stream=archive.read('ZUGFeRD-geprueft.pdf'),filetype='pdf') as pdf:
            self.assertEqual(pdf.embfile_get(0),xml)
        response=self.client.post('/v1/zugferd/validate',headers=self.headers,files={'file':('ordinary.pdf',self.pdf)})
        self.assertEqual(response.status_code,200,response.text[:200]);archive=zipfile.ZipFile(io.BytesIO(response.content));self.assertFalse(json.loads(archive.read('Ergebnis.json'))['compliant'])


if __name__=='__main__':unittest.main()
