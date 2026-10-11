import io
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pymupdf as fitz
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).parent))
from expansion_documents import svg_export, font_outlines, flatten_layers, editable_powerpoint, translate_original


class ExpandedDocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.source = self.directory/'input.pdf'
        with fitz.open() as pdf:
            page=pdf.new_page(width=500,height=400)
            page.draw_rect((20,20,480,380),fill=(.9,.95,1),color=(0,0,1))
            page.insert_textbox((40,40,450,100),'Original heading',fontsize=18)
            page.insert_textbox((40,130,240,210),'Left column content.\nSecond line remains.',fontsize=12)
            page.insert_textbox((270,130,450,210),'Right column content.',fontsize=12)
            page.insert_link({'kind':fitz.LINK_URI,'from':fitz.Rect(40,40,200,70),'uri':'https://example.org'})
            pdf.save(self.source)

    def tearDown(self):
        self.temp.cleanup()

    def test_svg_contains_paths_and_no_external_resources(self):
        svg_export(self.source,self.directory)
        with zipfile.ZipFile(self.directory/'output.zip') as archive:
            self.assertEqual(archive.namelist(),['Seite-0001.svg'])
            data=archive.read('Seite-0001.svg');root=ET.fromstring(data)
        self.assertGreater(len(root.findall('.//{http://www.w3.org/2000/svg}path')),5)
        self.assertEqual(len(root.findall('.//{http://www.w3.org/2000/svg}image')),0)
        self.assertEqual(root.get('viewBox'),'0 0 500 400')
        for node in root.iter():
            for key,value in node.attrib.items():
                if key.endswith('href'):self.assertTrue(value.startswith(('#','data:image/')))

    @unittest.skipUnless(shutil.which('gs'),'Ghostscript fehlt')
    def test_real_font_outlines_preserve_vector_picture(self):
        font_outlines(self.source,self.directory)
        with fitz.open(self.directory/'output.pdf') as pdf:
            self.assertEqual(len(pdf),1)
            self.assertEqual(pdf[0].get_fonts(),[])
            self.assertEqual(pdf[0].get_text().strip(),'')
            self.assertGreater(len(pdf[0].get_drawings()),4)
            self.assertEqual(len(pdf[0].get_images()),0)

    @unittest.skipUnless(shutil.which('gs'),'Ghostscript fehlt')
    def test_flatten_removes_hidden_layer_but_keeps_visible_text(self):
        with fitz.open() as pdf:
            page=pdf.new_page();visible=pdf.add_ocg('Visible',on=True);hidden=pdf.add_ocg('Hidden',on=False)
            page.insert_text((40,60),'VISIBLE LAYER',oc=visible)
            page.insert_text((40,120),'HIDDEN LAYER',oc=hidden)
            pdf.save(self.source)
        flatten_layers(self.source,self.directory)
        with fitz.open(self.directory/'output.pdf') as pdf:
            self.assertFalse(pdf.get_ocgs())
            self.assertIn('VISIBLE LAYER',pdf[0].get_text())
            self.assertNotIn('HIDDEN LAYER',pdf[0].get_text())

    def test_powerpoint_has_editable_styled_text(self):
        editable_powerpoint(self.source,self.directory)
        pptx=Presentation(self.directory/'output.pptx')
        self.assertEqual(len(pptx.slides),1)
        text=[shape for shape in pptx.slides[0].shapes if shape.has_text_frame]
        self.assertIn('Original heading','\n'.join(shape.text for shape in text))
        self.assertGreater(len(text),2)
        heading=next(shape for shape in text if shape.text=='Original heading')
        heading.text='Edited in PowerPoint'
        saved=io.BytesIO();pptx.save(saved)
        self.assertIn('Edited in PowerPoint','\n'.join(shape.text for shape in Presentation(io.BytesIO(saved.getvalue())).slides[0].shapes if shape.has_text_frame))

    def test_translation_replaces_text_without_deleting_background_or_links(self):
        with fitz.open(self.source) as pdf:
            old_drawings=len(pdf[0].get_drawings())
            translate_original(pdf,self.directory,lambda text:'New heading' if 'heading' in text else 'New text',8)
        with fitz.open(self.directory/'output.pdf') as pdf:
            self.assertEqual(len(pdf),1)
            text=pdf[0].get_text()
            self.assertIn('New heading',text)
            self.assertNotIn('Original',text)
            self.assertNotIn('column content',text)
            self.assertEqual(len(pdf[0].get_drawings()),old_drawings)
            self.assertEqual(pdf[0].get_links()[0]['uri'],'https://example.org')
            self.assertLess(pdf[0].search_for('New heading')[0].y0,80)

    def test_translation_overflow_creates_no_partial_output(self):
        calls=[]
        def translation(text):
            calls.append(text)
            return 'Fits' if len(calls)==1 else 'Unreasonably long translation. '*100
        with fitz.open(self.source) as pdf:
            original=pdf[0].get_text()
            with self.assertRaisesRegex(ValueError,'Seite 1, Absatz 2'):
                translate_original(pdf,self.directory,translation,8)
            self.assertEqual(pdf[0].get_text(),original)
        self.assertFalse((self.directory/'output.pdf').exists())

    def test_translation_html_is_escaped(self):
        with fitz.open(self.source) as pdf:
            translate_original(pdf,self.directory,lambda text:'<b>Text</b>',8)
        with fitz.open(self.directory/'output.pdf') as pdf:
            self.assertIn('<b>Text</b>',pdf[0].get_text())


if __name__=='__main__':unittest.main()
