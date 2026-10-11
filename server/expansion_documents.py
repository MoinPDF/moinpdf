"""Bounded vector, editable-slide and original-position export workers."""
from __future__ import annotations
import html
import io
import json
import math
import subprocess
import zipfile
from pathlib import Path

import pymupdf as fitz
from defusedxml import ElementTree as SafeET
from extended_documents import open_pdf, settings

MAX_OUTPUT = 80 * 1024 * 1024


def svg_export(source, directory):
    total = 0
    with open_pdf(source) as pdf, zipfile.ZipFile(directory / 'output.zip', 'w', zipfile.ZIP_DEFLATED) as archive:
        for number, page in enumerate(pdf, 1):
            data = page.get_svg_image(text_as_path=True).encode('utf-8')
            total += len(data)
            if len(data) > 20 * 1024 * 1024 or total > MAX_OUTPUT:
                raise ValueError('Der SVG-Export überschreitet 20 MB pro Seite oder 80 MB insgesamt. Bitte einen kleineren Seitenbereich exportieren.')
            root = SafeET.fromstring(data)
            for node in root.iter():
                if node.tag.split('}')[-1] in {'script', 'foreignObject'}:
                    raise ValueError('Unerwarteter aktiver Inhalt im SVG-Export.')
                for key, value in node.attrib.items():
                    if key.split('}')[-1] in {'href', 'src'} and not value.startswith(('#', 'data:image/')):
                        raise ValueError('Der SVG-Export enthält eine externe Referenz.')
            archive.writestr(f'Seite-{number:04}.svg', data)


def vector_rewrite(source, directory, outlines=False):
    with open_pdf(source) as original:
        pages = len(original)
    command = ['gs', '-q', '-dSAFER', '-dBATCH', '-dNOPAUSE', '-sDEVICE=pdfwrite', '-dAutoRotatePages=/None']
    if outlines:
        command += ['-dNoOutputFonts']
    command += ['-sOutputFile=' + str(directory / 'output.pdf'), str(source)]
    try:
        result = subprocess.run(command, capture_output=True, timeout=100)
    except (OSError, subprocess.TimeoutExpired) as error:
        raise ValueError('Die Vektorverarbeitung ist nicht verfügbar oder hat das Zeitlimit überschritten.') from error
    output = directory / 'output.pdf'
    if result.returncode or not output.exists() or output.stat().st_size > MAX_OUTPUT:
        raise ValueError('Die Vektorausgabe konnte nicht vollständig erzeugt werden (höchstens 80 MB).')
    with open_pdf(output) as checked:
        if len(checked) != pages:
            raise ValueError('Die Seitenzahl der Vektorausgabe stimmt nicht überein.')
        if outlines and any(page.get_fonts() for page in checked):
            raise ValueError('Nicht alle Schriften konnten in Konturen umgewandelt werden.')
        if not outlines and checked.get_ocgs():
            raise ValueError('Die Ebenen konnten nicht vollständig reduziert werden.')
    note = ('Schrift als Vektorkonturen exportiert; Text ist nicht mehr auswählbar oder durchsuchbar. ' if outlines else
            'Die anfänglich sichtbaren Ebenen wurden zu einer festen Ansicht zusammengeführt. Versteckte Ebeneninhalte sind entfernt. ')
    note += 'Das Ergebnis ist eine neue PDF. Formularfunktion, Tags, Links, Anhänge und digitale Signaturen werden nicht garantiert übernommen.'
    (directory / 'result-note.json').write_text(json.dumps({'note': note}))


def font_outlines(source, directory):
    vector_rewrite(source, directory, True)


def flatten_layers(source, directory):
    vector_rewrite(source, directory)


def editable_powerpoint(source, directory):
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE
    deck = Presentation()
    with open_pdf(source) as pdf:
        if len(pdf) > 100:
            raise ValueError('Für Folien bitte höchstens 100 PDF-Seiten verwenden.')
        deck.slide_width = Pt(pdf[0].rect.width)
        deck.slide_height = Pt(pdf[0].rect.height)
        total_characters = 0
        for number, page in enumerate(pdf, 1):
            blocks = page.get_text('dict', flags=fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES)['blocks']
            lines = [line for block in blocks if block['type'] == 0 for line in block['lines']]
            if any(abs(line['dir'][0] - 1) > .001 or abs(line['dir'][1]) > .001 for line in lines):
                raise ValueError(f'Seite {number} enthält gedrehten oder vertikalen Text. Dafür bitte den Seitenbild-Modus verwenden.')
            if page.rotation:
                raise ValueError(f'Seite {number} ist gedreht. Bitte zuerst die Seitendrehung normalisieren oder den Seitenbild-Modus verwenden.')
            total_characters += sum(len(span['text']) for line in lines for span in line['spans'])
            if total_characters > 1_000_000 or len(lines) > 5000:
                raise ValueError('Zu viele Textobjekte für bearbeitbare Folien.')
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            scale = min(deck.slide_width / Pt(page.rect.width), deck.slide_height / Pt(page.rect.height))
            left = (deck.slide_width - Pt(page.rect.width) * scale) / 2
            top = (deck.slide_height - Pt(page.rect.height) * scale) / 2
            # Remove only text from a separate background document; keep images
            # and vector artwork in the visible background, then add real runs.
            background = fitz.open(); background.insert_pdf(pdf, from_page=number-1, to_page=number-1)
            try:
                bg = background[0]
                for line in lines:
                    bg.add_redact_annot(fitz.Rect(line['bbox']), fill=False, cross_out=False)
                bg.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE, text=fitz.PDF_REDACT_TEXT_REMOVE)
                pix = bg.get_pixmap(matrix=fitz.Matrix(min(2,1800/max(bg.rect.width,bg.rect.height)), min(2,1800/max(bg.rect.width,bg.rect.height))),alpha=False)
                slide.shapes.add_picture(io.BytesIO(pix.tobytes('png')), int(left), int(top), width=int(Pt(page.rect.width)*scale), height=int(Pt(page.rect.height)*scale))
            finally:
                background.close()
            for line in lines:
                for span in line['spans']:
                    if not span['text'].strip():
                        continue
                    rect = fitz.Rect(span['bbox']); size = span['size']
                    box = slide.shapes.add_textbox(int(left+Pt(rect.x0)*scale), int(top+Pt(rect.y0)*scale), max(1,int(Pt(rect.width+size*.3)*scale)), max(1,int(Pt(rect.height+size*.3)*scale)))
                    frame = box.text_frame; frame.clear(); frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
                    frame.word_wrap = False; frame.auto_size = MSO_AUTO_SIZE.NONE
                    paragraph = frame.paragraphs[0]; paragraph.space_before = paragraph.space_after = Pt(0)
                    run = paragraph.add_run(); run.text = span['text']; run.font.size = Pt(size*scale)
                    run.font.name = span['font'].split('+')[-1]
                    run.font.bold = bool(span['flags'] & 16); run.font.italic = bool(span['flags'] & 2)
                    run.font.color.rgb = RGBColor.from_string(f"{span['color']:06X}")
    deck.save(directory / 'output.pptx')
    if (directory / 'output.pptx').stat().st_size > MAX_OUTPUT:
        raise ValueError('Die Folienausgabe überschreitet 80 MB.')


def translate_original(pdf, directory, translate, minimum_font_size=8):
    if isinstance(minimum_font_size,bool) or not isinstance(minimum_font_size,(float,int)) or not math.isfinite(minimum_font_size) or not 6 <= minimum_font_size <= 18:
        raise ValueError('Die Mindestschriftgröße muss zwischen 6 und 18 pt liegen.')
    plan = []; total = 0
    for number, page in enumerate(pdf, 1):
        blocks = [block for block in page.get_text('dict', flags=fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES)['blocks'] if block['type'] == 0]
        if not blocks:
            raise ValueError(f'Seite {number} enthält keinen auslesbaren Text. Bitte zuerst OCR verwenden.')
        rectangles = [fitz.Rect(block['bbox']) for block in blocks]
        for i, rect in enumerate(rectangles):
            if any((rect & other).get_area() > .5 for other in rectangles[:i]):
                raise ValueError(f'Seite {number} enthält überlappende Textabsätze. Bitte die Gegenüberstellung verwenden.')
        page_plan = []
        for index, (block, rect) in enumerate(zip(blocks, rectangles),1):
            if any(abs(line['dir'][0]-1)>.001 or abs(line['dir'][1])>.001 for line in block['lines']):
                raise ValueError(f'Seite {number}, Absatz {index}: gedrehter/vertikaler Text wird im Original-Layout nicht unterstützt.')
            spans = [span for line in block['lines'] for span in line['spans']]
            original = '\n'.join(''.join(span['text'] for span in line['spans']) for line in block['lines'])
            if not original.strip():
                continue
            total += len(original)
            if total > 60000:
                raise ValueError('Bitte höchstens 60.000 Textzeichen pro Übersetzung verwenden.')
            translated = translate(original)
            if not translated.strip() or len(translated)>120000:
                raise ValueError('Die Übersetzung enthält keinen Text oder ist zu lang.')
            dominant = max(spans, key=lambda span:len(span['text']))
            size = max(float(dominant['size']), minimum_font_size)
            css = f"* {{margin:0;padding:0}} body {{font-family:sans-serif;font-size:{size}pt;line-height:1;color:#{dominant['color']:06X}}}"
            markup = '<div>'+html.escape(translated).replace('\n','<br>')+'</div>'
            # Fit all blocks in isolated blank pages before modifying any input.
            probe = fitz.open(); probe_page = probe.new_page(width=max(page.mediabox.width,rect.x1+1),height=max(page.mediabox.height,rect.y1+1))
            try:
                spare, scale = probe_page.insert_htmlbox(rect,markup,css=css,scale_low=minimum_font_size/size)
                if spare<0:
                    raise ValueError(f'Seite {number}, Absatz {index}: Übersetzung passt bei mindestens {minimum_font_size:g} pt nicht in den vorhandenen Bereich. Bitte Gegenüberstellung wählen oder Text kürzen.')
            finally:
                probe.close()
            page_plan.append((rect,markup,css,minimum_font_size/size))
        plan.append(page_plan)
    for page, page_plan in zip(pdf,plan):
        links = page.get_links()
        for rect, *_ in page_plan:
            page.add_redact_annot(rect, fill=False, cross_out=False)
        page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE,graphics=fitz.PDF_REDACT_LINE_ART_NONE,text=fitz.PDF_REDACT_TEXT_REMOVE)
        for rect, markup, css, scale_low in page_plan:
            if page.insert_htmlbox(rect,markup,css=css,scale_low=scale_low)[0]<0:
                raise ValueError('Der vollständige Text konnte nicht platziert werden. Es wurde keine Ausgabe erstellt.')
        # Recreate links removed by redaction, keeping original destinations.
        remaining = {(link['kind'],tuple(link['from'])) for link in page.get_links()}
        for link in links:
            if (link['kind'],tuple(link['from'])) not in remaining:
                page.insert_link({key:value for key,value in link.items() if key not in {'xref','id'}})
    pdf.save(directory/'output.pdf',garbage=4,deflate=True)
