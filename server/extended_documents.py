"""Document helpers used by the isolated, time-limited conversion worker."""
from __future__ import annotations
import html
import io
import json
import math
import os
import re
import subprocess
import zipfile
from pathlib import Path

import pymupdf as fitz


def open_pdf(source):
    pdf=fitz.open(source)
    if pdf.needs_pass or not 0 < len(pdf) <= 200:
        pdf.close()
        raise ValueError('Bitte eine unverschlüsselte PDF mit 1 bis 200 Seiten verwenden.')
    return pdf


def settings(directory):
    path=directory/'settings.json'
    return json.loads(path.read_text()) if path.exists() else {}


def report(directory,title,items,note):
    (directory/'report.json').write_text(json.dumps({'title':title,'items':items,'note':note},ensure_ascii=False,indent=2))


def story_pdf(markup,output):
    story=fitz.Story(html=markup,user_css='body {font-family: sans-serif; font-size: 10pt} table {border-collapse:collapse; width:100%} td,th {padding:5pt; border:0.5pt solid #bbb} h1 {font-size:18pt}')
    writer=fitz.DocumentWriter(str(output))
    def rectangle(index,filled):
        if index>=300:raise ValueError('Das Ergebnis überschreitet 300 Seiten.')
        return fitz.Rect(0,0,595,842),fitz.Rect(42,42,553,800),None
    try:story.write(writer,rectangle)
    finally:writer.close()


def html_pdf(source,directory):
    from bs4 import BeautifulSoup
    if source.stat().st_size>5*1024*1024:raise ValueError('HTML-Dateien dürfen höchstens 5 MB groß sein.')
    soup=BeautifulSoup(source.read_bytes(),'html.parser')
    for node in soup(['script','style','link','iframe','object','embed','base','meta','input','button','video','audio','source','svg','img']):node.decompose()
    for node in soup.find_all(True):
        node.attrs={key:value for key,value in node.attrs.items() if key in {'colspan','rowspan'}}
    story_pdf(str(soup),directory/'output.pdf')


def powerpoint(source,directory):
    from pptx import Presentation
    from pptx.util import Inches
    deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
    with open_pdf(source) as pdf:
        if len(pdf)>100:raise ValueError('Für Folien bitte höchstens 100 PDF-Seiten verwenden.')
        for page in pdf:
            pix=page.get_pixmap(matrix=fitz.Matrix(min(2,1800/max(page.rect.width,page.rect.height)),min(2,1800/max(page.rect.width,page.rect.height))),alpha=False)
            slide=deck.slides.add_slide(deck.slide_layouts[6]);scale=min(deck.slide_width/pix.width,deck.slide_height/pix.height)
            width,height=int(pix.width*scale),int(pix.height*scale)
            slide.shapes.add_picture(io.BytesIO(pix.tobytes('png')),int((deck.slide_width-width)/2),int((deck.slide_height-height)/2),width=width,height=height)
    deck.save(directory/'output.pptx')


def orient(source,directory):
    options=settings(directory);threshold=float(options.get('confidence',15))
    if not 5<=threshold<=30:raise ValueError('Ungültige Erkennungsschwelle.')
    items=[]
    with open_pdf(source) as pdf:
        if len(pdf)>50:raise ValueError('Bitte höchstens 50 Seiten pro Ausrichtungsauftrag verwenden.')
        for index,page in enumerate(pdf):
            image=directory/'orientation.png';page.get_pixmap(dpi=150,alpha=False).save(image)
            try:
                found=subprocess.run(['tesseract',str(image),'stdout','--psm','0'],capture_output=True,text=True,timeout=20)
                rotation=re.search(r'Rotate:\s*(\d+)',found.stdout);confidence=re.search(r'Orientation confidence:\s*([\d.]+)',found.stdout)
                angle=int(rotation.group(1)) if rotation else 0;certainty=float(confidence.group(1)) if confidence else 0
                changed=found.returncode==0 and certainty>=threshold and angle in {90,180,270}
                if changed:page.set_rotation((page.rotation+angle)%360)
                items.append({'page':index+1,'status':'gedreht' if changed else 'unverändert','rotation':angle if changed else 0,'confidence':certainty,'detail':'Bitte prüfen' if certainty<threshold else 'Ausrichtung erkannt'})
            except subprocess.TimeoutExpired:
                items.append({'page':index+1,'status':'unverändert','detail':'Zeitlimit der Erkennung erreicht; bitte manuell prüfen.'})
        pdf.save(directory/'output.pdf',garbage=4,deflate=True)
    report(directory,'Automatische Seitenausrichtung',items,'Unsichere Seiten bleiben unverändert. Bitte den Bericht und die gedrehten Seiten prüfen.')


def form_detect(source,directory):
    import cv2
    import numpy as np
    items=[]
    with open_pdf(source) as pdf:
        if len(pdf)>50:raise ValueError('Bitte höchstens 50 Seiten pro Formularerkennung verwenden.')
        names={widget.field_name for page in pdf for widget in page.widgets() or []}
        for index,page in enumerate(pdf):
            candidates=[];existing=[fitz.Rect(widget.rect) for widget in page.widgets() or []]
            for drawing in page.get_drawings():
                for item in drawing['items']:
                    if item[0]=='re':candidates.append(fitz.Rect(item[1]))
                    elif item[0]=='l':
                        first,last=item[1:3]
                        if abs(first.y-last.y)<1 and 35<abs(first.x-last.x)<page.rect.width*.9:
                            candidates.append(fitz.Rect(min(first.x,last.x),first.y-16,max(first.x,last.x),first.y))
            # Rectangular outlines on scanned pages.
            if not candidates:
                pix=page.get_pixmap(dpi=120,colorspace=fitz.csGRAY,alpha=False)
                pixels=np.frombuffer(pix.samples,dtype=np.uint8).reshape(pix.height,pix.width)
                contours,_=cv2.findContours(cv2.threshold(pixels,170,255,cv2.THRESH_BINARY_INV)[1],cv2.RETR_LIST,cv2.CHAIN_APPROX_SIMPLE)
                for contour in contours:
                    polygon=cv2.approxPolyDP(contour,cv2.arcLength(contour,True)*.025,True)
                    if len(polygon)==4:
                        x,y,w,h=cv2.boundingRect(polygon)
                        candidates.append(fitz.Rect(x,y,x+w,y+h)*fitz.Matrix(72/120,72/120)*page.derotation_matrix)
            accepted=[]
            for rectangle in candidates:
                w,h=rectangle.width,rectangle.height
                if not (8<=h<=85 and 8<=w<=500 and w*h<page.mediabox.width*page.mediabox.height*.15):continue
                if rectangle.is_empty or not page.mediabox.contains(rectangle):continue
                if page.get_textbox(rectangle).strip():continue
                if any((rectangle & other).get_area() > min(rectangle.get_area(),other.get_area())*.6 for other in existing+accepted):continue
                number=len(items)+1
                if number>500:raise ValueError('Mehr als 500 mögliche Felder. Bitte einen kleineren Seitenbereich verwenden.')
                name=f'Erkannt_{number}'
                while name in names:name+='_'
                names.add(name);widget=fitz.Widget();widget.field_name=name
                checkbox=.7<=w/h<=1.3 and h<=24
                widget.field_type=fitz.PDF_WIDGET_TYPE_CHECKBOX if checkbox else fitz.PDF_WIDGET_TYPE_TEXT
                widget.rect=rectangle;widget.text_font='Helv';widget.text_fontsize=10;widget.border_width=0
                page.add_widget(widget);accepted.append(rectangle)
                items.append({'page':index+1,'name':name,'type':'Checkbox' if checkbox else 'Textfeld','rectangle':list(rectangle)})
        if not items:raise ValueError('Keine leeren Kästchen oder Eingabelinien erkannt. Nutze „Formular erstellen“ für eigene Felder.')
        pdf.save(directory/'output.pdf',garbage=4,deflate=True)
    report(directory,'Erkannte Formularfelder',items,'Automatische Vorschläge: Namen, Feldtypen und Positionen im Werkzeug „Formular erstellen“ prüfen.')


def print_check(source,directory):
    minimum=float(settings(directory).get('dpi',150));items=[]
    if not 72<=minimum<=600:raise ValueError('Bitte eine Mindestauflösung zwischen 72 und 600 dpi wählen.')
    with open_pdf(source) as pdf:
        for number,page in enumerate(pdf,1):
            issues=[];images=[]
            for image in page.get_image_info():
                matrix=fitz.Matrix(image['transform']);width=math.hypot(matrix.a,matrix.b);height=math.hypot(matrix.c,matrix.d)
                dpi=min(image['width']*72/max(width,.001),image['height']*72/max(height,.001))
                images.append(round(dpi))
                if dpi<minimum:issues.append(f'Bildauflösung {dpi:.0f} dpi liegt unter {minimum:.0f} dpi.')
            fonts=[]
            for font in page.get_fonts(full=True):
                embedded=bool(pdf.extract_font(font[0])[3]) if font[0] else False
                if not embedded:fonts.append(font[3])
            if fonts:issues.append('Schriften ohne eingebettete Fontdatei: '+', '.join(sorted(set(fonts))))
            items.append({'page':number,'status':'Hinweise' if issues else 'Kein Befund in diesen Prüfungen','size_mm':[round(page.rect.width*25.4/72,1),round(page.rect.height*25.4/72,1)],'image_dpi':images,'trimbox':list(page.trimbox),'bleedbox':list(page.bleedbox),'findings':issues})
    report(directory,'Druckqualität: Basisprüfung',items,'Geprüft: Bildauflösung, Font-Einbettung und Seiten-/Beschnittboxen. Keine PDF/X-Zertifizierung; Farbseparation, Überdruck, Transparenzen und Druckerei-Vorgaben zusätzlich prüfen.')


def accessibility(source,directory):
    items=[]
    with open_pdf(source) as pdf:
        catalog=pdf.pdf_catalog();language=pdf.xref_get_key(catalog,'Lang');tags=pdf.xref_get_key(catalog,'StructTreeRoot')
        items.append({'scope':'Dokument','title':pdf.metadata.get('title',''),'language':language[1] if language[0]!='null' else '', 'structure_tree_present':tags[0]!='null','bookmarks':len(pdf.get_toc()),'findings':[label for ok,label in [(bool(pdf.metadata.get('title')),'Dokumenttitel fehlt.'),(language[0]!='null','Dokumentsprache fehlt.'),(tags[0]!='null','Strukturbaum (Tags) fehlt.')] if not ok]})
        for number,page in enumerate(pdf,1):
            widgets=list(page.widgets() or []);issues=[]
            if not page.get_text().strip():issues.append('Kein auslesbarer Text; bei Textseiten OCR nötig.')
            for widget in widgets:
                if not widget.field_label:issues.append(f'Formularfeld „{widget.field_name}“ ohne erklärende Beschriftung (Tooltip).')
            items.append({'page':number,'findings':issues,'text_characters':len(page.get_text()),'images':len(page.get_image_info()),'form_fields':len(widgets)})
    report(directory,'Barrierefreiheit: Basisprüfung',items,'Dieser Check weist technische Lücken aus. Lesereihenfolge, Alternativtexte, Kontrast und Bedienbarkeit müssen zusätzlich geprüft werden. Kein Nachweis einer PDF/UA- oder WCAG-Konformität.')


def redaction_check(source,directory):
    options=settings(directory);terms=[t.strip() for t in options.get('terms','').splitlines() if t.strip()]
    if len(terms)>50 or sum(map(len,terms))>5000:raise ValueError('Bitte höchstens 50 kurze Suchbegriffe eingeben.')
    items=[]
    with open_pdf(source) as pdf:
        for number,page in enumerate(pdf,1):
            for term in terms:
                for rectangle in page.search_for(term):items.append({'page':number,'status':'Text auslesbar','text':term,'rectangle':list(rectangle)})
            drawings=page.get_drawings()
            for trace in page.get_texttrace():
                text=''.join(chr(character[0]) for character in trace['chars'] if 0<=character[0]<=0x10ffff)
                if not text.strip():continue
                rectangle=fitz.Rect(trace['bbox'])
                # Later opaque dark vectors overlapping text can be cosmetic redactions.
                for drawing in drawings:
                    fill=drawing.get('fill')
                    if fill is None or max(fill)>.15 or drawing.get('fill_opacity',1)<.95 or drawing.get('seqno',-1)<=trace.get('seqno',0):continue
                    overlap=rectangle & drawing['rect']
                    if not overlap.is_empty and overlap.get_area()>rectangle.get_area()*.35:
                        hidden=page.get_textbox(overlap).strip()
                        items.append({'page':number,'status':'Verdächtige Überdeckung','text':hidden or text,'rectangle':list(overlap)})
                        break
        names=pdf.embfile_names()
        if names:items.append({'scope':'Anhänge','status':'Zusätzlich prüfen','files':names})
        metadata='\n'.join(str(v or '') for v in pdf.metadata.values())+'\n'+pdf.get_xml_metadata()
        for term in terms:
            if term.casefold() in metadata.casefold():items.append({'scope':'Metadaten','status':'Text auslesbar','text':term})
    report(directory,'Schwärzungen prüfen',items,'Suchbegriffe, Metadaten und mögliche dunkle Vektor-Überdeckungen geprüft. Keine Freigabe für sichere Schwärzung: Bilder, maskierte Inhalte, alte PDF-Revisionen und Anlagen zusätzlich kontrollieren. Ein leerer Bericht beweist keine vollständige Entfernung.')


def curve_image(pixels,amount,side):
    import cv2
    import numpy as np
    height,width=pixels.shape[:2];padding=int(abs(amount))+2
    x,y=np.meshgrid(np.arange(width,dtype=np.float32),np.arange(height+2*padding,dtype=np.float32))
    position=x/max(1,width-1)
    bend=(1-position)**2 if side=='left' else position**2
    return cv2.remap(pixels,x,y-padding+amount*bend,cv2.INTER_LINEAR,borderMode=cv2.BORDER_CONSTANT,borderValue=(255,255,255))


def dewarp(source,directory):
    import cv2
    import numpy as np
    options=settings(directory);percent=float(options.get('curve',5));side=options.get('side','left')
    if not -20<=percent<=20 or side not in {'left','right'}:raise ValueError('Ungültige Buchfalz-Korrektur.')
    output=fitz.open()
    try:
        with open_pdf(source) as pdf:
            for page in pdf:
                scale=min(2,2200/max(page.rect.width,page.rect.height));pix=page.get_pixmap(matrix=fitz.Matrix(scale,scale),colorspace=fitz.csRGB,alpha=False)
                pixels=np.frombuffer(pix.samples,dtype=np.uint8).reshape(pix.height,pix.width,3)
                corrected=curve_image(pixels,pix.height*percent/100,side)
                encoded=cv2.imencode('.png',cv2.cvtColor(corrected,cv2.COLOR_RGB2BGR))[1].tobytes()
                target=output.new_page(width=page.rect.width,height=corrected.shape[0]/scale);target.insert_image(target.rect,stream=encoded)
        output.save(directory/'output.pdf',deflate=True)
    finally:output.close()


def translate(source,directory):
    import requests
    options=settings(directory);target=options.get('target','en');language=options.get('language','auto')
    if not re.fullmatch(r'[a-z]{2,3}',target) or not re.fullmatch(r'auto|[a-z]{2,3}',language):raise ValueError('Ungültige Sprache.')
    endpoint=os.environ.get('MOINPDF_TRANSLATE_URL','').rstrip('/')
    if not endpoint:raise ValueError('Auf dem Server ist noch kein eigener LibreTranslate-Dienst eingerichtet.')
    rows=[];total=0
    session=requests.Session();session.trust_env=False
    try:
        with open_pdf(source) as pdf:
            for number,page in enumerate(pdf,1):
                text=page.get_text(sort=True).strip();total+=len(text)
                if total>60000:raise ValueError('Bitte höchstens 60.000 Textzeichen pro Übersetzung verwenden.')
                if not text:raise ValueError(f'Seite {number} enthält keinen auslesbaren Text. Bitte zuerst OCR verwenden.')
                # Short paired rows remain readable and can flow across page boundaries.
                for start in range(0,len(text),1500):
                    chunk=text[start:start+1500]
                    payload={'q':chunk,'source':language,'target':target,'format':'text'}
                    if os.environ.get('MOINPDF_TRANSLATE_KEY'):payload['api_key']=os.environ['MOINPDF_TRANSLATE_KEY']
                    response=session.post(endpoint+'/translate',json=payload,timeout=30,allow_redirects=False)
                    if response.status_code!=200:raise ValueError('Der Übersetzungsdienst hat den Auftrag nicht angenommen. Sprachen und Dienstkonfiguration prüfen.')
                    translated=response.json().get('translatedText')
                    if not isinstance(translated,str):raise ValueError('Der Übersetzungsdienst hat keinen übersetzten Text geliefert.')
                    rows.append(f'<tr><td><b>Seite {number}</b><p>{html.escape(chunk)}</p></td><td>{html.escape(translated)}</td></tr>')
        story_pdf('<h1>Original und Übersetzung</h1><p>Maschinelle Übersetzung – bitte prüfen.</p><table><tr><th>Original</th><th>Übersetzung</th></tr>'+''.join(rows)+'</table>',directory/'output.pdf')
    finally:session.close()


def word_merge(source,directory):
    from openpyxl import load_workbook
    from lxml import etree
    spreadsheet=directory/'data.xlsx'
    with zipfile.ZipFile(spreadsheet) as archive:
        if sum(item.file_size for item in archive.infolist())>40*1024*1024:raise ValueError('Die entpackte Excel-Datei ist zu groß.')
    workbook=load_workbook(spreadsheet,read_only=True,data_only=False)
    try:
        sheet=workbook.worksheets[0]
        if sheet.max_row>51 or sheet.max_column>50:raise ValueError('Bitte eine Tabelle mit höchstens 50 Datenzeilen und 50 Spalten verwenden.')
        rows=list(sheet.iter_rows(max_row=51,max_col=50))
        headers=[str(cell.value).strip() if cell.value is not None else '' for cell in rows[0]]
        while headers and not headers[-1]:headers.pop()
        if not headers or len(headers)>50 or len(set(headers))!=len(headers) or any(not name for name in headers):raise ValueError('Die erste Tabellenzeile braucht eindeutige, nicht leere Spaltennamen (höchstens 50).')
        records=[]
        for row in rows[1:]:
            if not any(cell.value is not None for cell in row):continue
            if len(records)>=50:raise ValueError('Bitte höchstens 50 Datensätze pro Serienbrief verwenden.')
            if any(cell.data_type=='f' for cell in row):raise ValueError('Bitte Excel-Formeln vor dem Import durch ihre berechneten Werte ersetzen.')
            records.append({name:'' if cell.value is None else str(cell.value) for name,cell in zip(headers,row)})
        if not records:raise ValueError('Die Tabelle enthält keine Datensätze.')
    finally:workbook.close()
    with zipfile.ZipFile(source) as archive:
        if sum(item.file_size for item in archive.infolist())>40*1024*1024:raise ValueError('Die entpackte Word-Vorlage ist zu groß.')
        if any(Path(item.filename).is_absolute() or '..' in Path(item.filename).parts or '\\' in item.filename for item in archive.infolist()):raise ValueError('Ungültige Word-Paketstruktur.')
        package={item.filename:archive.read(item) for item in archive.infolist()}
    namespace={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    parser=etree.XMLParser(resolve_entities=False,no_network=True)
    pattern=re.compile(r'\{\{\s*([^{}]+?)\s*\}\}')
    paths=[name for name in package if re.fullmatch(r'word/(document|header\d+|footer\d+)\.xml',name)]
    found=set()
    for path in paths:
        root=etree.fromstring(package[path],parser)
        for paragraph in root.xpath('//w:p',namespaces=namespace):found.update(pattern.findall(''.join(paragraph.xpath('.//w:t/text()',namespaces=namespace))))
    if not found:raise ValueError('Keine Platzhalter gefunden. Verwende in Word z. B. {{Name}} passend zum Excel-Spaltennamen.')
    if found-set(headers):raise ValueError('Excel-Spalten fehlen: '+', '.join(sorted(found-set(headers))))
    output=directory/'merged';output.mkdir()
    for index,record in enumerate(records,1):
        updated=dict(package)
        for path in paths:
            root=etree.fromstring(package[path],parser)
            for paragraph in root.xpath('//w:p',namespaces=namespace):
                nodes=paragraph.xpath('.//w:t',namespaces=namespace);combined=''.join(node.text or '' for node in nodes)
                for match in reversed(list(pattern.finditer(combined))):
                    offset=0;positions=[]
                    for node in nodes:
                        text=node.text or '';positions.append((node,offset,offset+len(text)));offset+=len(text)
                    for node,start,end in positions:
                        if end<=match.start() or start>=match.end():continue
                        text=node.text or '';prefix=text[:max(0,match.start()-start)];suffix=text[max(0,match.end()-start):]
                        node.text=prefix+(record[match.group(1)] if start<=match.start()<end else '')+suffix
                        node.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
                # Other runs keep their original formatting.
            updated[path]=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
        # Do not resolve linked documents/images while converting the template.
        for path in [p for p in updated if p.endswith('.rels')]:
            root=etree.fromstring(updated[path],parser)
            for child in list(root):
                if child.get('TargetMode')=='External':root.remove(child)
            updated[path]=etree.tostring(root,xml_declaration=True,encoding='UTF-8')
        with zipfile.ZipFile(output/f'Serienbrief-{index:03}.docx','w',zipfile.ZIP_DEFLATED) as archive:
            for path,data in updated.items():archive.writestr(path,data)
    report(directory,'Word-Serienbriefe',[{'number':i+1,'filename':f'Serienbrief-{i+1:03}.pdf'} for i in range(len(records))],'Ein PDF je nicht leerer Tabellenzeile. Erste Tabelle, Platzhalter {{Spaltenname}}, höchstens 50 Datensätze. Formeln vorher in Werte umwandeln.')
