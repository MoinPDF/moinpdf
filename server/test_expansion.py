import io
import os
import struct
import unittest
import zipfile
from email.message import EmailMessage
from unittest.mock import patch

import pymupdf
from docx import Document
from extract_msg.ole_writer import OleWriter
from fastapi.testclient import TestClient
from server.app import app
from server.test_service import pdf_bytes


def msg_bytes():
    writer = OleWriter()
    writer.addEntry('__properties_version1.0', b'\0'*32 + struct.pack('<IIII', 0x340D0003, 0, 0x40000, 0))
    for prop, value in [('001A','IPM.Note'), ('0037','MSG Grüße'), ('1000','Nachricht aus Outlook.'), ('0C1A','Testsender'), ('0C1F','test@example.org'), ('0E04','Empfang')]:
        writer.addEntry('__substg1.0_'+prop+'001F', value.encode('utf-16-le'))
    for stream in ('00020102','00030102','00040102'):
        writer.addEntry('__nameid_version1.0/__substg1.0_'+stream, b'')
    prefix = '__attach_version1.0_#00000000/'
    writer.addEntry(prefix+'__properties_version1.0', b'\0'*8 + struct.pack('<IIII',0x37050003,0,1,0))
    writer.addEntry(prefix+'__substg1.0_3707001F', 'Anlage.txt'.encode('utf-16-le'))
    writer.addEntry(prefix+'__substg1.0_37010102', b'MSG attachment bytes')
    stream=io.BytesIO();writer.write(stream)
    return stream.getvalue()


class ExpansionTests(unittest.TestCase):
    def setUp(self):
        self.environment = patch.dict(os.environ, {'MOINPDF_API_KEY':'synthetic-test-'+'x'*40})
        self.environment.start()
        self.client = TestClient(app)
        self.headers = {'Authorization':'Bearer '+os.environ['MOINPDF_API_KEY']}

    def tearDown(self):
        self.client.close();self.environment.stop()

    def post(self, path, name, data, values=None):
        return self.client.post(path, headers=self.headers, files={'file':(name,data)}, data=values or {})

    def test_word_text_and_scan_rejection(self):
        response = self.post('/v1/pdf-to-word','test.pdf',pdf_bytes())
        self.assertEqual(response.status_code,200,response.text[:200])
        document = Document(io.BytesIO(response.content))
        self.assertIn('MoinPDF test','\n'.join(p.text for p in document.paragraphs))
        blank=pymupdf.open();blank.new_page()
        response=self.post('/v1/pdf-to-word','scan.pdf',blank.tobytes());blank.close()
        self.assertEqual(response.status_code,422)
        self.assertIn('OCR',response.json()['detail'])

    def test_word_mixed_native_and_scanned_pages_are_editable(self):
        native = pymupdf.open(stream=pdf_bytes(), filetype='pdf')
        scan_source = pymupdf.open()
        page = scan_source.new_page()
        page.insert_text((70, 120), 'SCANNED INVOICE 4815', fontsize=22)
        page.insert_text((70, 160), 'Total amount 1234.56 EUR', fontsize=18)
        image = page.get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes('png')
        scan = native.new_page()
        scan.insert_image(scan.rect, stream=image)
        response = self.post('/v1/pdf-to-word', 'mixed.pdf', native.tobytes())
        native.close(); scan_source.close()
        self.assertEqual(response.status_code, 200, response.text[:200])
        document = Document(io.BytesIO(response.content))
        text = '\n'.join(p.text for p in document.paragraphs)
        self.assertIn('MoinPDF test', text)
        self.assertIn('SCANNED INVOICE 4815', text)
        self.assertIn('1234.56', text)
        self.assertIn('OCR', response.headers.get('x-moinpdf-note', ''))

    def test_eml_preserves_unicode_and_duplicate_attachments(self):
        message=EmailMessage();message['From']='Test <test@example.org>';message['To']='empfang@example.org';message['Subject']='Grüße aus Köln'
        message.set_content('Moin Welt! Prüftext mit Umlauten.')
        for payload in (b'first\x00\xff',b'second\x01'):
            message.add_attachment(payload,maintype='application',subtype='octet-stream',filename='../Anlage.bin')
        response=self.post('/v1/email','test.eml',message.as_bytes())
        self.assertEqual(response.status_code,200,response.text[:200])
        archive=zipfile.ZipFile(io.BytesIO(response.content))
        self.assertEqual(archive.read('Anlagen/_Anlage.bin'),b'first\x00\xff')
        self.assertEqual(archive.read('Anlagen/_Anlage-2.bin'),b'second\x01')
        with pymupdf.open(stream=archive.read('E-Mail.pdf'),filetype='pdf') as pdf:
            text='\n'.join(page.get_text() for page in pdf)
        self.assertIn('Grüße aus Köln',text);self.assertIn('Prüftext',text)

    def test_word_preserves_table_and_image(self):
        from PIL import Image
        document=pymupdf.open();page=document.new_page()
        for row,values in enumerate([['Artikel','Preis'],['Buch','12.50']]):
            for column,value in enumerate(values):
                left=40+column*200;top=60+row*35
                page.draw_rect((left,top,left+200,top+35))
                page.insert_text((left+8,top+22),value,fontsize=12)
        image=io.BytesIO();Image.new('RGB',(80,80),'orange').save(image,format='PNG')
        page.insert_image(pymupdf.Rect(40,180,140,280),stream=image.getvalue())
        response=self.post('/v1/pdf-to-word','layout.pdf',document.tobytes());document.close()
        self.assertEqual(response.status_code,200,response.text[:200])
        word=Document(io.BytesIO(response.content))
        self.assertGreaterEqual(len(word.inline_shapes),1)
        self.assertIn('12.50',' '.join(cell.text for table in word.tables for row in table.rows for cell in row.cells))

    def test_html_email_and_pdf_only(self):
        message=EmailMessage();message['Subject']='HTML Nachricht';message.set_content('<html><script>DO_NOT_RUN</script><p>Lesbarer Text</p><img src="https://invalid.example/image.png"></html>',subtype='html')
        response=self.post('/v1/email','html.eml',message.as_bytes(),{'include_attachments':'false'})
        self.assertEqual(response.status_code,200,response.text[:200])
        with pymupdf.open(stream=response.content,filetype='pdf') as pdf:text='\n'.join(p.get_text() for p in pdf)
        self.assertIn('Lesbarer Text',text);self.assertNotIn('DO_NOT_RUN',text)

    def test_msg_with_attachment(self):
        response=self.post('/v1/email','test.msg',msg_bytes())
        self.assertEqual(response.status_code,200,response.text[:200])
        archive=zipfile.ZipFile(io.BytesIO(response.content))
        self.assertEqual(archive.read('Anlagen/Anlage.txt'),b'MSG attachment bytes')
        with pymupdf.open(stream=archive.read('E-Mail.pdf'),filetype='pdf') as pdf:text='\n'.join(p.get_text() for p in pdf)
        self.assertIn('MSG Grüße',text);self.assertIn('Nachricht aus Outlook.',text)

    def test_invalid_types(self):
        self.assertEqual(self.post('/v1/email','file.exe',b'bad').status_code,422)
        self.assertEqual(self.post('/v1/email','file.eml',b'bad').status_code,422)
        self.assertEqual(self.post('/v1/pdf-to-word','file.pdf',b'bad').status_code,422)


if __name__=='__main__':unittest.main()
