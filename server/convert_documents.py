"""Bounded conversion workers, launched in the service's disposable process group."""
from __future__ import annotations
import json
import re
import sys
from email import policy
from email.parser import BytesParser
from pathlib import Path


def clean_text(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value or ''))


def html_text(value):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(value, 'html.parser')
    for node in soup(['script', 'style', 'iframe', 'object', 'head']):
        node.decompose()
    return soup.get_text('\n', strip=True)


def word(source: Path, directory: Path):
    import csv
    import subprocess
    import pymupdf
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.shared import Pt
    from pdf2docx import Converter
    document = Document()
    ocr_count = 0
    empty_count = 0
    any_text = False
    converter = None
    try:
        with pymupdf.open(source) as pdf:
            if pdf.needs_pass or not 0 < len(pdf) <= 200:
                raise ValueError('Bitte eine unverschlüsselte PDF mit höchstens 200 Seiten verwenden.')
            native = [i for i, page in enumerate(pdf) if page.get_text().strip()]
            if native:
                converter = Converter(str(source))
                settings = {**converter.default_settings, 'ignore_page_error': False, 'raw_exceptions': True, 'multi_processing': False}
                converter.parse(pages=native, **settings)
            for index, page in enumerate(pdf):
                if index in native:
                    converter.pages[index].make_docx(document)
                    any_text = True
                    continue
                # OCR pages become real paragraphs, never a page-sized image
                # with an invisible text layer. Native pages keep pdf2docx layout.
                ocr_count += 1
                image = directory / 'word-ocr.png'
                scale = min(3.0, 3200 / max(page.rect.width, page.rect.height))
                page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False).save(image)
                try:
                    completed = subprocess.run(['tesseract', str(image), 'stdout', '-l', 'deu+eng', '--psm', '3', 'tsv'], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=120)
                    lines = {}
                    for item in csv.DictReader(completed.stdout.decode('utf-8').splitlines(), delimiter='\t', quoting=csv.QUOTE_NONE):
                        text = clean_text(item.get('text', '')).strip()
                        if item.get('level') == '5' and text:
                            key = (item['block_num'], item['par_num'], item['line_num'])
                            lines.setdefault(key, []).append(text)
                except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as error:
                    raise ValueError(f'OCR auf Seite {index + 1} fehlgeschlagen. Bitte Scanqualität prüfen.') from error
                finally:
                    image.unlink(missing_ok=True)
                section = document.add_section(WD_SECTION.NEW_PAGE) if index else document.sections[0]
                section.page_width, section.page_height = Pt(page.rect.width), Pt(page.rect.height)
                section.left_margin = section.right_margin = Pt(36)
                section.top_margin = section.bottom_margin = Pt(36)
                for words in lines.values():
                    document.add_paragraph(' '.join(words))
                    any_text = True
                if not lines:
                    empty_count += 1
                    document.add_paragraph('')
        if not any_text:
            raise ValueError('Auch per OCR wurde kein Text erkannt. Bitte die Scanqualität prüfen.')
        document.save(directory / 'output.docx')
        note = f'{ocr_count} Seiten per OCR in bearbeitbare Absätze umgewandelt. Scanlayout und erkannte Texte bitte kontrollieren.' if ocr_count else 'Text, Tabellen und Bilder mit Layoutrekonstruktion übernommen. Ausgabe bitte kontrollieren.'
        if empty_count:
            note += f' {empty_count} Seiten ohne erkannten Text.'
        (directory / 'word-note.txt').write_text(note, encoding='utf-8')
    finally:
        if converter:
            converter.close()


def parse_email(source: Path):
    if source.suffix == '.eml':
        message = BytesParser(policy=policy.default).parsebytes(source.read_bytes())
        if not any(message.get(name) for name in ('From', 'To', 'Subject', 'Date')):
            raise ValueError('Die Datei enthält keine erkennbaren E-Mail-Kopfzeilen.')
        headers = [(label, message.get(name, '')) for label, name in [('Von', 'From'), ('An', 'To'), ('Cc', 'Cc'), ('Datum', 'Date'), ('Betreff', 'Subject')]]
        part = message.get_body(preferencelist=('plain', 'html'))
        body = part.get_content() if part else ''
        if part and part.get_content_type() == 'text/html':
            body = html_text(body)
        attachments = []
        # Include inline resources as well as ordinary attachments; do not fetch remote URLs.
        body_ids = set()
        if part:
            body_ids.add(id(part))
        def collect(node):
            if id(node) in body_ids:
                return
            if node.get_content_type() == 'message/rfc822':
                payload = node.get_payload()
                data = b'\n'.join(item.as_bytes(policy=policy.default) for item in payload) if isinstance(payload, list) else node.get_payload(decode=True)
                attachments.append((node.get_filename() or 'Nachricht.eml', data or b''))
            elif node.is_multipart():
                for item in node.iter_parts():
                    collect(item)
            elif node.get_filename() or node.get_content_disposition() == 'attachment' or node.get_content_maintype() not in ('text', 'multipart'):
                attachments.append((node.get_filename() or f'Anlage-{len(attachments)+1}', node.get_payload(decode=True) or b''))
        collect(message)
        return headers, body, attachments
    import extract_msg
    message = extract_msg.Message(str(source))
    try:
        headers = [('Von', message.sender), ('An', message.to), ('Cc', message.cc), ('Datum', message.date), ('Betreff', message.subject)]
        body = message.body or html_text(message.htmlBody or b'')
        attachments = []
        for item in message.attachments:
            data = item.data
            if hasattr(data, 'exportBytes'):
                data = data.exportBytes()
            if not isinstance(data, bytes):
                raise ValueError('Eine MSG-Anlage kann nicht vollständig gelesen werden. Bitte die Nachricht als EML exportieren.')
            attachments.append((item.name or f'Anlage-{len(attachments)+1}', data))
        return headers, body, attachments
    finally:
        message.close()


def email_pdf(source: Path, directory: Path):
    from docx import Document
    from docx.shared import Pt, Cm
    headers, body, attachments = parse_email(source)
    if len(body) > 2_000_000 or len(attachments) > 100 or sum(len(data) for _, data in attachments) > 80 * 1024 * 1024:
        raise ValueError('Die E-Mail überschreitet das Limit von 100 Anlagen oder 80 MB Anlageninhalt.')
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    document.styles['Normal'].font.name = 'Liberation Sans'
    document.styles['Normal'].font.size = Pt(10)
    document.add_heading(clean_text(dict(headers).get('Betreff')) or 'E-Mail', 0)
    for label, value in headers:
        if value:
            paragraph = document.add_paragraph()
            paragraph.add_run(label + ': ').bold = True
            paragraph.add_run(clean_text(value))
    document.add_paragraph('')
    for line in clean_text(body).splitlines() or ['(Nachricht ohne Textinhalt)']:
        for offset in range(0, max(1, len(line)), 4000):
            document.add_paragraph(line[offset:offset+4000])
    folder = directory / 'attachments'
    folder.mkdir()
    names = []
    used = set()
    for index, (name, data) in enumerate(attachments):
        name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(name)).strip(' .')[:140] or f'Anlage-{index+1}'
        stem, suffix = Path(name).stem, Path(name).suffix
        candidate, number = name, 1
        while candidate.casefold() in used:
            number += 1
            candidate = f'{stem}-{number}{suffix}'
        used.add(candidate.casefold())
        (folder / candidate).write_bytes(data)
        names.append(candidate)
    if names:
        document.add_heading('Anlagen', 1)
        for name in names:
            document.add_paragraph(name)
    document.save(directory / 'email.docx')
    (directory / 'attachments.json').write_text(json.dumps(names), encoding='utf-8')


if __name__ == '__main__':
    directory = Path(sys.argv[3])
    try:
        import extended_documents as extra
        import recovery_documents as recovery
        import expansion_documents as expansion
        import tag_documents as tags
        actions={'word':word,'email':email_pdf,'powerpoint':extra.powerpoint,'html':extra.html_pdf,'orient':extra.orient,'form':extra.form_detect,'print':extra.print_check,'accessibility':extra.accessibility,'redaction':extra.redaction_check,'dewarp':extra.dewarp,'translate':extra.translate,'word-merge':extra.word_merge}
        actions.update({'optimize': recovery.optimize, 'ebook': recovery.ebook, 'svg':expansion.svg_export, 'font-outlines':expansion.font_outlines, 'flatten-layers':expansion.flatten_layers, 'tag-structure':tags.inspect_tags, 'tag-repair':tags.repair_tags})
        actions[sys.argv[1]](Path(sys.argv[2]), directory)
    except ValueError as error:
        (directory / 'conversion-error.json').write_text(json.dumps({'detail': str(error)}))
        sys.exit(1)
